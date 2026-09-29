"""Тесты экспортёров результата (txt/docx/json/srt)."""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from audio_transcriber.domain.enums import ExportFormat
from audio_transcriber.domain.models import Speaker, TranscriptEntry, TranscriptionResult
from audio_transcriber.export.docx_exporter import DocxExporter
from audio_transcriber.export.factory import create_exporter
from audio_transcriber.export.json_exporter import JsonExporter
from audio_transcriber.export.srt_exporter import SrtExporter
from audio_transcriber.export.txt_exporter import TxtExporter


@pytest.fixture
def sample_result(tmp_path: Path) -> TranscriptionResult:
    ivan = Speaker(id="SPEAKER_00", display_name="Иван")
    maria = Speaker(id="SPEAKER_01", display_name="Мария")
    return TranscriptionResult(
        source_path=tmp_path / "call.mp3",
        language="ru",
        duration=5.0,
        entries=[
            TranscriptEntry(start=0.0, end=1.5, text="Привет, как дела?", speaker=ivan),
            TranscriptEntry(start=1.5, end=3.0, text="Хорошо, спасибо!", speaker=maria),
            TranscriptEntry(start=3.0, end=4.0, text="Без определённого говорящего"),
        ],
        speakers=[ivan, maria],
    )


def test_txt_exporter_writes_lines_with_speakers(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    output_path = tmp_path / "out.txt"

    TxtExporter().export(sample_result, output_path)

    content = output_path.read_text(encoding="utf-8")
    assert "Иван: Привет, как дела?" in content
    assert "Мария: Хорошо, спасибо!" in content
    assert "?: Без определённого говорящего" in content


def test_json_exporter_round_trips_data(sample_result: TranscriptionResult, tmp_path: Path) -> None:
    output_path = tmp_path / "out.json"

    JsonExporter().export(sample_result, output_path)

    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert payload["language"] == "ru"
    assert len(payload["entries"]) == 3
    assert payload["entries"][0]["speaker"] == "SPEAKER_00"
    assert payload["entries"][2]["speaker"] is None
    assert {speaker["display_name"] for speaker in payload["speakers"]} == {"Иван", "Мария"}


def test_srt_exporter_writes_valid_blocks(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    output_path = tmp_path / "out.srt"

    SrtExporter().export(sample_result, output_path)

    content = output_path.read_text(encoding="utf-8")
    assert "00:00:00,000 --> 00:00:01,500" in content
    assert "Иван: Привет, как дела?" in content


def test_docx_exporter_creates_readable_document(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    output_path = tmp_path / "out.docx"

    DocxExporter().export(sample_result, output_path)

    document = Document(str(output_path))
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Иван" in full_text
    assert "Хорошо, спасибо!" in full_text


@pytest.mark.parametrize(
    ("export_format", "expected_type"),
    [
        (ExportFormat.TXT, TxtExporter),
        (ExportFormat.DOCX, DocxExporter),
        (ExportFormat.JSON, JsonExporter),
        (ExportFormat.SRT, SrtExporter),
    ],
)
def test_factory_creates_matching_exporter(
    export_format: ExportFormat, expected_type: type
) -> None:
    assert isinstance(create_exporter(export_format), expected_type)


_SUMMARY = "Тема: созвон\nЧто сделать:\n- подготовить отчёт — Мария"


def test_txt_exporter_puts_summary_before_participants_and_entries(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    result = replace(sample_result, participants=["Иван", "Мария"], summary=_SUMMARY)
    output_path = tmp_path / "out.txt"

    TxtExporter().export(result, output_path)

    content = output_path.read_text(encoding="utf-8")
    assert content.startswith("Резюме встречи")
    assert "Тема: созвон" in content
    assert content.index("Резюме встречи") < content.index("Участники:")
    assert content.index("Резюме встречи") < content.index("Иван: Привет")


def test_txt_exporter_without_summary_has_no_block(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    output_path = tmp_path / "out.txt"

    TxtExporter().export(sample_result, output_path)

    assert "Резюме встречи" not in output_path.read_text(encoding="utf-8")


def test_docx_exporter_puts_summary_at_top(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    result = replace(sample_result, summary=_SUMMARY)
    output_path = tmp_path / "out.docx"

    DocxExporter().export(result, output_path)

    document = Document(str(output_path))
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Резюме встречи" in full_text
    assert full_text.index("Резюме встречи") < full_text.index("Привет, как дела?")


def test_json_exporter_includes_summary_field(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    result = replace(sample_result, summary=_SUMMARY)
    output_path = tmp_path / "out.json"

    JsonExporter().export(result, output_path)

    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert payload["summary"] == _SUMMARY


def test_json_exporter_summary_is_null_without_llm(
    sample_result: TranscriptionResult, tmp_path: Path
) -> None:
    output_path = tmp_path / "out.json"

    JsonExporter().export(sample_result, output_path)

    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert payload["summary"] is None
