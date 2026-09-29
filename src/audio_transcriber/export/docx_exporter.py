"""Экспорт стенограммы в документ Microsoft Word (.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from audio_transcriber.domain.models import TranscriptionResult
from audio_transcriber.export.annotations import entry_markers
from audio_transcriber.export.timestamps import format_timestamp
from audio_transcriber.utils.exceptions import ExportError


class DocxExporter:
    """Реализует протокол ``ResultExporter`` для формата DOCX."""

    def export(self, result: TranscriptionResult, output_path: Path) -> None:
        document = Document()
        document.add_heading(result.source_path.name, level=1)

        if result.summary:
            document.add_heading("Резюме встречи", level=2)
            for line in result.summary.splitlines():
                if line.strip():
                    document.add_paragraph(line)

        if result.participants:
            document.add_heading("Участники", level=2)
            for participant in result.participants:
                document.add_paragraph(participant, style="List Bullet")

        threshold = result.low_confidence_threshold
        for entry in result.entries:
            speaker_label = entry.speaker.display_name if entry.speaker else "?"
            paragraph = document.add_paragraph()
            paragraph.add_run(f"[{format_timestamp(entry.start)}] {speaker_label}: ").bold = True
            paragraph.add_run(entry.text + entry_markers(entry, threshold))

        try:
            document.save(str(output_path))
        except OSError as exc:
            raise ExportError(f"Не удалось сохранить файл {output_path}: {exc}") from exc
